"""Export the rewritten quant deck to a spreadsheet.

    python scripts/export_quant_xlsx.py

Writes data/GRE_Quant_Notes_rewritten.xlsx with one row per concept:
Day | Group | Concept | Explanation | Example | Watch out.

This is a one-way export for reading, printing or pasting into your own
workbook - content/quant/*.md stays the source of truth, so edit the markdown
and re-run this rather than the other way round.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "data" / "quant.json"
OUT = ROOT / "data" / "GRE_Quant_Notes_rewritten.xlsx"

HEADERS = ["Day", "Group", "Concept", "Explanation", "Example", "Watch out"]
WIDTHS = [6, 32, 46, 90, 90, 60]


def main() -> None:
    if not DECK.exists():
        raise SystemExit(f"{DECK} is missing - run scripts/build_data.py first.")
    deck = json.loads(DECK.read_text(encoding="utf-8"))

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quant Notes"
    sheet.append(HEADERS)

    for day, group in enumerate(deck["groups"], start=1):
        for item in group["items"]:
            blocks = {b["label"]: b["text"] for b in item["blocks"]}
            explanation = "\n\n".join(
                text for label, text in blocks.items() if label.startswith("Explanation")
            )
            example = "\n\n".join(
                text for label, text in blocks.items() if label.startswith("Example")
            )
            sheet.append(
                [day, group["title"], item["label"], explanation, example, blocks.get("Watch out", "")]
            )

    header_fill = PatternFill("solid", fgColor="1F6FEB")
    for column, (header, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = sheet.cell(row=1, column=column)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{sheet.max_row}"

    workbook.save(OUT)
    print(f"{sheet.max_row - 1} concepts -> {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
